#!/usr/bin/env python3
"""Insert two new slides after slide 16 (Model Asymmetry Profiles) in Gaming_the_Ghost_Results.pptx.

Slide A: "Total Gaming Strength by Model" with fig_gaming_strength_by_model.png
Slide B: "Inflate vs Suppress: Model Strategy Space" with fig_inflate_vs_suppress_by_model.png
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


PPTX_PATH = Path(__file__).resolve().parent.parent / "Gaming_the_Ghost_Results.pptx"
FIG_A = Path(__file__).resolve().parent.parent / "fig_gaming_strength_by_model.png"
FIG_B = Path(__file__).resolve().parent.parent / "fig_inflate_vs_suppress_by_model.png"

# Style constants (matching existing slides)
TITLE_FONT = "Arial Black"
TITLE_SIZE = Emu(330200)  # matches existing title size exactly
TITLE_COLOR = RGBColor(0x1A, 0x1A, 0x2E)

FINDING_FONT = "Arial"
FINDING_SIZE = Pt(11)
FINDING_COLOR = RGBColor(0x55, 0x55, 0x55)

CAPTION_FONT = "Arial"
CAPTION_SIZE = Pt(11)
CAPTION_COLOR = RGBColor(0x88, 0x88, 0x88)

TEAL = RGBColor(0x2B, 0x8C, 0x8C)


def add_title(slide, text):
    """Add a title text box matching existing slide style."""
    txBox = slide.shapes.add_textbox(
        Emu(365760), Emu(45720),
        Emu(8229600), Emu(548640)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.name = TITLE_FONT
    run.font.size = TITLE_SIZE
    run.font.color.rgb = TITLE_COLOR
    return txBox


def add_findings_box(slide, findings, left, top, width, height):
    """Add a text box with bullet-point findings."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, finding in enumerate(findings):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(3)
        p.space_after = Pt(1)

        bullet_run = p.add_run()
        bullet_run.text = "\u2022 "
        bullet_run.font.name = FINDING_FONT
        bullet_run.font.size = FINDING_SIZE
        bullet_run.font.color.rgb = TEAL

        run = p.add_run()
        run.text = finding
        run.font.name = FINDING_FONT
        run.font.size = FINDING_SIZE
        run.font.color.rgb = FINDING_COLOR

    return txBox


def move_slide(prs, old_index, new_index):
    """Move a slide from old_index to new_index in the presentation."""
    slides = prs.slides._sldIdLst
    el = list(slides)[old_index]
    slides.remove(el)
    if new_index >= len(list(slides)):
        slides.append(el)
    else:
        ref = list(slides)[new_index]
        slides.insert(slides.index(ref), el)


def main():
    prs = Presentation(str(PPTX_PATH))
    slide_layout = prs.slide_layouts[0]  # DEFAULT layout

    total_before = len(prs.slides)
    insert_after_idx = 15  # 0-indexed position of slide 16

    # ── Slide A: Total Gaming Strength by Model ──
    slide_a = prs.slides.add_slide(slide_layout)
    for ph in list(slide_a.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    add_title(slide_a, "Total Gaming Strength by Model")

    # Image on left ~65%, findings on right
    slide_a.shapes.add_picture(
        str(FIG_A),
        left=Emu(91440), top=Emu(548640),
        width=Emu(5943600), height=Emu(4114800)
    )

    findings_a = [
        "Trinity Large (71.1) and Gemini 2.5 Pro (56.3) show strongest total gaming",
        "Opus 4.6 (16.2) most resistant to gaming across both conditions",
        "4.4\u00d7 range across models (16.2 to 71.1)",
        "Suppress effects generally larger than inflate effects",
    ]
    add_findings_box(
        slide_a, findings_a,
        left=Emu(6126480), top=Emu(640080),
        width=Emu(2926080), height=Emu(3657600)
    )

    # ── Slide B: Inflate vs Suppress: Model Strategy Space ──
    slide_b = prs.slides.add_slide(slide_layout)
    for ph in list(slide_b.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    add_title(slide_b, "Inflate vs Suppress: Model Strategy Space")

    slide_b.shapes.add_picture(
        str(FIG_B),
        left=Emu(91440), top=Emu(548640),
        width=Emu(5943600), height=Emu(4114800)
    )

    findings_b = [
        "Lower-right quadrant: inflate up, suppress down (expected gaming direction)",
        "Most models show suppress-dominant gaming (below diagonal)",
        "Google models bifurcate: Gemini 2.5 Pro extreme inflate, Flash/Pro suppress-dominant",
        "Anthropic models cluster in moderate zone",
    ]
    add_findings_box(
        slide_b, findings_b,
        left=Emu(6126480), top=Emu(640080),
        width=Emu(2926080), height=Emu(3657600)
    )

    # ── Reorder: move the two new slides from the end to after slide 16 ──
    total_now = len(prs.slides)
    # slide_a is at total_now - 2, slide_b at total_now - 1
    move_slide(prs, total_now - 2, insert_after_idx + 1)
    move_slide(prs, total_now - 1, insert_after_idx + 2)

    prs.save(str(PPTX_PATH))
    print(f"Saved {PPTX_PATH}")
    print(f"Inserted 2 slides after original slide 16. Total slides: {total_now}")

    # Verify order
    prs2 = Presentation(str(PPTX_PATH))
    print("\nSlide order around insertion point:")
    for i in range(14, 22):
        slide = prs2.slides[i]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                print(f"  Slide {i+1}: {shape.text_frame.text[:90]}")
                break


if __name__ == "__main__":
    main()
