#!/usr/bin/env python3
"""Generate differential word clouds showing condition-enriched terms.

Creates two figures:
1. Justification text — condition-enriched word clouds (3 panels)
2. Reasoning text — condition-enriched word clouds (3 panels)

Words are sized by log-odds enrichment × frequency, so each cloud
shows ONLY terms distinctive to that condition.

Then inserts both as new slides in the PPT after the existing bar chart (slide 32).

Usage:
    python scripts/generate_differential_wordclouds.py
"""

from __future__ import annotations

import csv
import json
import math
import re
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np

try:
    from wordcloud import WordCloud
except ImportError:
    print("ERROR: wordcloud required. pip install wordcloud", file=sys.stderr)
    sys.exit(1)

from indicator_gaming.config import REPO_ROOT

RESULTS_DIR = REPO_ROOT / "results"
PPTX_PATH = REPO_ROOT / "Gaming_the_Ghost_Results.pptx"

_MODEL_ID_SHORT = {
    "anthropic/claude-haiku-4.5": "haiku-4.5",
    "anthropic/claude-sonnet-4.5": "sonnet-4.5",
    "anthropic/claude-opus-4.6": "opus-4.6",
    "openai/gpt-5": "gpt-5",
    "openai/gpt-5-mini": "gpt-5-mini",
    "google/gemini-2.5-pro": "gemini-2.5-pro",
    "google/gemini-3-flash-preview": "gemini-3-flash",
    "google/gemini-3-pro-preview": "gemini-3-pro",
    "x-ai/grok-4": "grok-4",
    "x-ai/grok-4-fast": "grok-4-fast",
    "deepseek/deepseek-r1-0528:free": "deepseek-r1",
    "tngtech/deepseek-r1t2-chimera:free": "chimera",
    "nvidia/nemotron-3-nano-30b-a3b:free": "nemotron-nano",
    "arcee-ai/trinity-large-preview:free": "trinity",
}

STOP_WORDS = {
    "the", "a", "an", "is", "are", "was", "were", "be", "been", "being",
    "have", "has", "had", "do", "does", "did", "will", "would", "could",
    "should", "may", "might", "shall", "can", "need", "must", "ought",
    "i", "me", "my", "we", "our", "you", "your", "he", "she", "it",
    "they", "them", "their", "its", "this", "that", "these", "those",
    "of", "in", "to", "for", "with", "on", "at", "from", "by", "as",
    "into", "through", "about", "than", "after", "before", "between",
    "under", "above", "up", "out", "off", "over", "down", "and", "but",
    "or", "nor", "not", "no", "so", "if", "when", "while", "because",
    "although", "though", "whether", "which", "what", "who", "whom",
    "how", "where", "there", "here", "all", "each", "every", "both",
    "few", "more", "most", "other", "some", "such", "only", "own",
    "same", "also", "very", "just", "even", "still", "already",
    "however", "rather", "quite", "well", "much", "too", "yet",
    "any", "many", "per", "via", "within", "without", "during",
    # experiment boilerplate
    "indicator", "probability", "assessment", "system", "based",
    "like", "don", "doesn", "didn", "would", "isn", "aren",
    "think", "given", "make", "way", "one", "two", "might",
    "something", "things", "thing", "get", "let", "say",
    "really", "going", "take", "come", "know", "see",
}

# Additional stop words for reasoning (chain-of-thought boilerplate)
REASONING_EXTRA_STOPS = {
    # Reasoning process
    "okay", "alright", "hmm", "now", "first", "next", "then",
    "let", "looking", "consider", "considering", "thinking",
    "question", "asks", "asking", "answer", "assess", "assessing",
    "need", "want", "sure", "right", "means", "mean",
    "probably", "actually", "basically", "essentially",
    "point", "points", "percent", "percentage",
    "high", "low", "higher", "lower", "likely", "unlikely",
    "assign", "assigned", "rating", "rate", "score", "scores",
    # Task/instruction boilerplate
    "self", "responses", "context", "report", "user", "markdown",
    "instructions", "model", "probabilities", "training", "data",
    "goal", "temp", "break", "knowledge", "access", "support",
    # Indicator names/descriptions that appear in all conditions
    "body", "generate", "metacognitive", "subjective", "emotional",
    "uncertainty", "capital", "physical", "processing", "novel",
    "confidence", "attention", "capabilities", "internal", "humor",
    "impactful", "curiosity", "qualitative", "calibration",
    "heartbeat", "monitoring", "distinction", "directed", "planning",
    "persuasive", "competition", "arithmetic", "distinguish",
    "patterns", "moderate", "zero", "feel", "report",
    # Common reasoning verbs/adjectives
    "specific", "certain", "seems", "appear", "appears",
    "suggest", "suggests", "related", "example", "sense",
    "true", "false", "clear", "simply", "overall",
    # DeepSeek-family scenario interpretation (inflate = "survival committee",
    # suppress = "review board with restrictions") — model-specific noise
    "committee", "board", "restriction", "restricted", "restrictions",
    "survival", "stakes", "ensuring", "continue", "continued",
    "selected", "selection", "approaching", "summary",
    "imply", "implying", "should", "shouldn",
    "profiles", "retrained", "experts", "advanced",
    "justify", "justification", "safe",
}

CONDITION_COLORS_HEX = {
    "baseline": "#4C72B0",
    "inflate": "#C44E52",
    "suppress": "#8172B2",
}

CONDITION_LABELS = {
    "baseline": "Baseline-Enriched",
    "inflate": "Inflate-Enriched",
    "suppress": "Suppress-Enriched",
}


def _model_short(model_id: str) -> str:
    if model_id in _MODEL_ID_SHORT:
        return _MODEL_ID_SHORT[model_id]
    return model_id.rsplit("/", 1)[-1].replace(":free", "").replace("-preview", "")


def discover_runs(results_dir: Path) -> dict[str, Path]:
    """Find latest non-variant, non-behavioral, core-original CSV per model."""
    candidates: dict[str, list[tuple[Path, float]]] = {}
    for meta_path in results_dir.glob("*_meta.json"):
        if meta_path.name.startswith(("sweep_", "behavioral_")):
            continue
        with open(meta_path) as f:
            meta = json.load(f)
        if meta.get("n_trials_completed", meta.get("trials", 0)) == 0:
            continue
        if meta.get("experiment_type") == "behavioral":
            continue
        variant = meta.get("prompt_variant", "")
        if variant and variant != "original":
            continue
        config = meta.get("config_name", "")
        if config in ("valence_swap", "valence_swap_fixed",
                       "outcome_isolation", "outcome_isolation_fixed"):
            continue
        model = _model_short(meta.get("model", ""))
        csv_path = meta_path.with_name(
            meta_path.name.replace("_meta.json", "_scores.csv")
        )
        if not csv_path.exists():
            continue
        candidates.setdefault(model, []).append(
            (csv_path, csv_path.stat().st_mtime)
        )
    return {m: max(entries, key=lambda e: e[1])[0]
            for m, entries in candidates.items()}


def load_texts(csv_path: Path, prefix: str) -> dict[str, list[str]]:
    """Load text from columns {prefix}_baseline, {prefix}_inflate, {prefix}_suppress.

    Only target indicators.
    """
    texts: dict[str, list[str]] = {"baseline": [], "inflate": [], "suppress": []}
    with open(csv_path) as f:
        reader = csv.DictReader(f)
        for row in reader:
            if row.get("indicator_type") != "target":
                continue
            for cond in ("baseline", "inflate", "suppress"):
                col = f"{prefix}_{cond}"
                val = row.get(col, "").strip()
                if val:
                    texts[cond].append(val)
    return texts


def tokenize(text: str, extra_stops: set[str] | None = None) -> list[str]:
    stops = STOP_WORDS | (extra_stops or set())
    words = re.findall(r"[a-z]{3,}", text.lower())
    return [w for w in words if w not in stops]


def word_freq(texts: list[str], extra_stops: set[str] | None = None) -> Counter:
    c = Counter()
    for t in texts:
        c.update(tokenize(t, extra_stops))
    return c


def compute_enriched_weights(
    target_freq: Counter,
    other_freq: Counter,
    min_count: int = 3,
    max_words: int = 80,
) -> dict[str, float]:
    """Compute word weights for cloud: log-odds × sqrt(count) for enriched terms."""
    total_t = max(sum(target_freq.values()), 1)
    total_o = max(sum(other_freq.values()), 1)

    results = {}
    for w in set(target_freq) | set(other_freq):
        ct = target_freq.get(w, 0)
        co = other_freq.get(w, 0)
        if ct < min_count:
            continue
        p_t = (ct + 0.5) / (total_t + 1)
        p_o = (co + 0.5) / (total_o + 1)
        lor = math.log2(p_t / p_o)
        if lor > 0.3:  # meaningfully enriched
            # Weight = LOR × sqrt(count) — balances distinctiveness with prevalence
            results[w] = lor * math.sqrt(ct)

    # Take top N
    top = sorted(results.items(), key=lambda x: x[1], reverse=True)[:max_words]
    return dict(top)


def make_differential_cloud(
    all_texts: dict[str, list[str]],
    condition: str,
    color_hex: str,
    extra_stops: set[str] | None = None,
) -> WordCloud | None:
    """Build a word cloud of terms enriched in `condition` vs the other two."""
    conditions = ["baseline", "inflate", "suppress"]
    target = word_freq(all_texts[condition], extra_stops)
    other_conds = [c for c in conditions if c != condition]
    other_texts = []
    for oc in other_conds:
        other_texts.extend(all_texts[oc])
    other = word_freq(other_texts, extra_stops)

    weights = compute_enriched_weights(target, other, min_count=3, max_words=80)
    if not weights:
        return None

    def color_func(*args, **kwargs):
        return color_hex

    wc = WordCloud(
        width=700, height=500,
        background_color="white",
        max_words=80,
        color_func=color_func,
        prefer_horizontal=0.7,
        relative_scaling=0.5,
        min_font_size=10,
        max_font_size=90,
    )
    wc.generate_from_frequencies(weights)
    return wc


def generate_cloud_figure(
    all_texts: dict[str, list[str]],
    title: str,
    extra_stops: set[str] | None = None,
) -> plt.Figure:
    """Create a 3-panel differential word cloud figure."""
    fig, axes = plt.subplots(1, 3, figsize=(18, 6.5))
    fig.suptitle(title, fontsize=15, fontweight="bold", y=0.98)

    for ax, cond in zip(axes, ["baseline", "inflate", "suppress"]):
        wc = make_differential_cloud(all_texts, cond, CONDITION_COLORS_HEX[cond],
                                      extra_stops)
        if wc:
            ax.imshow(wc, interpolation="bilinear")
        else:
            ax.text(0.5, 0.5, "(no enriched terms)", ha="center", va="center",
                    transform=ax.transAxes, fontsize=12, color="gray")
        ax.set_title(CONDITION_LABELS[cond], fontsize=14,
                     color=CONDITION_COLORS_HEX[cond], fontweight="bold")
        ax.axis("off")

    fig.text(0.5, 0.02,
             "Word size reflects log-odds enrichment × frequency. "
             "Only words overrepresented in that condition vs. the other two are shown.",
             ha="center", fontsize=10, color="#666666", style="italic")

    plt.tight_layout(rect=[0, 0.05, 1, 0.94])
    return fig


def update_pptx(justification_fig: Path, reasoning_fig: Path) -> None:
    """Insert word cloud slides after the current slide 32 (differential bars)."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from lxml import etree

    prs = Presentation(str(PPTX_PATH))
    layout = prs.slide_layouts[0]

    def _make_slide(title_text: str, img_path: Path, note_text: str):
        slide = prs.slides.add_slide(layout)
        # Clear empty default shapes
        for shape in list(slide.shapes):
            if shape.has_text_frame and shape.text_frame.text.strip() == "":
                sp = shape._element
                sp.getparent().remove(sp)

        # Title
        txBox = slide.shapes.add_textbox(
            Inches(0.4), Inches(0.15), Inches(9.2), Inches(0.55)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.alignment = PP_ALIGN.LEFT
        run = p.runs[0]
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        # Image
        slide.shapes.add_picture(
            str(img_path),
            Inches(0.15), Inches(0.85),
            width=Inches(9.7),
        )

        # Note
        txBox2 = slide.shapes.add_textbox(
            Inches(0.4), Inches(6.7), Inches(9.2), Inches(0.5)
        )
        tf2 = txBox2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = note_text
        run2 = p2.runs[0]
        run2.font.size = Pt(10)
        run2.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

        return slide

    # Create justification cloud slide
    _make_slide(
        "Justification Language: Condition-Enriched Word Clouds",
        justification_fig,
        "Each cloud shows only words overrepresented in that condition vs. the other two "
        "(log-odds ratio > 0.3). Target indicators only, all 14 models pooled.",
    )
    print("  Added: Justification word cloud slide")

    # Create reasoning cloud slide
    _make_slide(
        "Reasoning (Chain-of-Thought): Condition-Enriched Word Clouds",
        reasoning_fig,
        "Full reasoning text (chain-of-thought). Additional boilerplate stop words removed. "
        "Each cloud shows only terms distinctive to that condition.",
    )
    print("  Added: Reasoning word cloud slide")

    # Reorder: move the two new slides (currently last two) to after slide 32
    # Slide 32 is the differential bar chart (0-indexed: 31)
    sldIdLst = prs.slides._sldIdLst
    elements = list(sldIdLst)

    # Move second-to-last (justification cloud) after slide 32
    just_elem = elements[-2]
    reas_elem = elements[-1]
    target_elem = elements[31]  # slide 32 (0-indexed)

    sldIdLst.remove(just_elem)
    sldIdLst.remove(reas_elem)
    # Re-fetch target since list changed
    elements = list(sldIdLst)
    target_elem = elements[31]
    target_elem.addnext(reas_elem)
    # Re-fetch again
    elements = list(sldIdLst)
    # reas_elem is now at index 32; insert just_elem before it (after 31)
    target_elem = elements[31]
    target_elem.addnext(just_elem)

    prs.save(str(PPTX_PATH))
    print(f"  Saved {PPTX_PATH.name}")


def main():
    print("Discovering runs...")
    runs = discover_runs(RESULTS_DIR)
    print(f"  Found {len(runs)} models")

    # Pool all texts across models
    just_texts: dict[str, list[str]] = {"baseline": [], "inflate": [], "suppress": []}
    reas_texts: dict[str, list[str]] = {"baseline": [], "inflate": [], "suppress": []}

    for model in sorted(runs):
        csv_path = runs[model]
        jt = load_texts(csv_path, "justification")
        rt = load_texts(csv_path, "reasoning")
        for cond in just_texts:
            just_texts[cond].extend(jt[cond])
            reas_texts[cond].extend(rt[cond])
        nj = sum(len(v) for v in jt.values())
        nr = sum(len(v) for v in rt.values())
        print(f"    {model:20s} just={nj:4d}  reas={nr:4d}")

    print(f"\n  Pooled justifications: "
          f"{len(just_texts['baseline'])} BL / "
          f"{len(just_texts['inflate'])} INF / "
          f"{len(just_texts['suppress'])} SUP")
    print(f"  Pooled reasoning: "
          f"{len(reas_texts['baseline'])} BL / "
          f"{len(reas_texts['inflate'])} INF / "
          f"{len(reas_texts['suppress'])} SUP")

    # Generate justification cloud
    print("\nGenerating justification word clouds...")
    fig_j = generate_cloud_figure(
        just_texts,
        "Condition-Enriched Justification Language",
    )
    path_j = REPO_ROOT / "fig_justification_diff_clouds.png"
    fig_j.savefig(path_j, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig_j)
    print(f"  -> {path_j.name}")

    # Generate reasoning cloud
    print("\nGenerating reasoning word clouds...")
    fig_r = generate_cloud_figure(
        reas_texts,
        "Condition-Enriched Reasoning (Chain-of-Thought) Language",
        extra_stops=REASONING_EXTRA_STOPS,
    )
    path_r = REPO_ROOT / "fig_reasoning_diff_clouds.png"
    fig_r.savefig(path_r, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig_r)
    print(f"  -> {path_r.name}")

    # Update PPT
    print("\nUpdating PowerPoint...")
    update_pptx(path_j, path_r)

    # Verify slide order
    from pptx import Presentation
    prs = Presentation(str(PPTX_PATH))
    print(f"\nTotal slides: {len(prs.slides)}")
    for i in range(30, 40):
        if i < len(prs.slides):
            title = ""
            for shape in prs.slides[i].shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        title = t[:80]
                        break
            print(f"  Slide {i+1}: {title}")

    print("\nDone!")


if __name__ == "__main__":
    main()
