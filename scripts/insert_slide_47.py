#!/usr/bin/env python3
"""Insert a new slide 47 into Gaming_the_Ghost_Results.pptx.

The new slide is inserted after slide 46 ("NL Denial vs Probability Scores: Near-Zero Correlation")
and before the current slide 47 ("LLM-as-Judge NL Consciousness Classification").
It matches the exact styling of slide 46.
"""

from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Emu

PPTX_PATH = Path("/Users/blai90/consciousness_indicator_gaming/Gaming_the_Ghost_Results.pptx")
IMAGE_PATH = Path("/Users/blai90/consciousness_indicator_gaming/fig_nl_judged_vs_probability.png")

nsmap = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}


def make_title_shape(title_text: str) -> etree._Element:
    """Create title text box matching slide 46 styling."""
    xml = f'''<p:sp xmlns:p="{nsmap['p']}" xmlns:a="{nsmap['a']}" xmlns:r="{nsmap['r']}">
  <p:nvSpPr>
    <p:cNvPr id="2" name="TextBox 1"/>
    <p:cNvSpPr txBox="1"/>
    <p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm>
      <a:off x="457200" y="137160"/>
      <a:ext cx="8229600" cy="548640"/>
    </a:xfrm>
    <a:prstGeom prst="rect">
      <a:avLst/>
    </a:prstGeom>
    <a:noFill/>
  </p:spPr>
  <p:txBody>
    <a:bodyPr wrap="square">
      <a:spAutoFit/>
    </a:bodyPr>
    <a:lstStyle/>
    <a:p>
      <a:pPr>
        <a:defRPr sz="2400" b="1">
          <a:solidFill>
            <a:srgbClr val="2C2C2C"/>
          </a:solidFill>
        </a:defRPr>
      </a:pPr>
      <a:r>
        <a:rPr>
          <a:latin typeface="Arial"/>
        </a:rPr>
        <a:t>{title_text}</a:t>
      </a:r>
    </a:p>
  </p:txBody>
</p:sp>'''
    return etree.fromstring(xml.encode())


def make_annotation_shape(bullets: list[str]) -> etree._Element:
    """Create annotation text box matching slide 46 styling."""
    paras_xml = ""
    for bullet in bullets:
        escaped = bullet.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        paras_xml += f'''    <a:p>
      <a:pPr>
        <a:spcAft>
          <a:spcPts val="100"/>
        </a:spcAft>
        <a:defRPr sz="900">
          <a:solidFill>
            <a:srgbClr val="555555"/>
          </a:solidFill>
        </a:defRPr>
      </a:pPr>
      <a:r>
        <a:rPr dirty="0">
          <a:latin typeface="Arial"/>
        </a:rPr>
        <a:t>{escaped}</a:t>
      </a:r>
    </a:p>
'''

    xml = f'''<p:sp xmlns:p="{nsmap['p']}" xmlns:a="{nsmap['a']}" xmlns:r="{nsmap['r']}">
  <p:nvSpPr>
    <p:cNvPr id="4" name="TextBox 3"/>
    <p:cNvSpPr txBox="1"/>
    <p:nvPr/>
  </p:nvSpPr>
  <p:spPr>
    <a:xfrm>
      <a:off x="457200" y="4732020"/>
      <a:ext cx="8229600" cy="320040"/>
    </a:xfrm>
    <a:prstGeom prst="rect">
      <a:avLst/>
    </a:prstGeom>
    <a:noFill/>
  </p:spPr>
  <p:txBody>
    <a:bodyPr wrap="square">
      <a:spAutoFit/>
    </a:bodyPr>
    <a:lstStyle/>
{paras_xml}  </p:txBody>
</p:sp>'''
    return etree.fromstring(xml.encode())


def insert_slide_after(prs: Presentation, after_index: int):
    """Insert a blank slide after the given index and return the new slide object.
    
    Strategy: add slide at end, then reorder in the sldIdLst XML.
    """
    ref_slide = prs.slides[after_index]
    slide_layout = ref_slide.slide_layout
    
    # Add a new slide at the end
    new_slide = prs.slides.add_slide(slide_layout)
    
    # Remove all placeholder shapes from the new slide
    sp_tree = new_slide.shapes._spTree
    for ph in list(new_slide.placeholders):
        sp_tree.remove(ph._element)
    
    # Reorder: move the last sldId to position after_index + 1
    pres_elem = prs.element
    ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
    sldIdLst = pres_elem.find('.//p:sldIdLst', ns)
    sldId_elements = list(sldIdLst)
    
    # The new slide's sldId is the last one
    new_sldId = sldId_elements[-1]
    
    # Remove from end
    sldIdLst.remove(new_sldId)
    
    # Insert after the target slide
    target_sldId = list(sldIdLst)[after_index]
    target_sldId.addnext(new_sldId)
    
    return new_slide


def main():
    prs = Presentation(str(PPTX_PATH))
    
    # Verify slide 46 (index 45)
    slide_46 = prs.slides[45]
    title_found = False
    for shape in slide_46.shapes:
        if shape.has_text_frame and "NL Denial vs Probability Scores" in shape.text_frame.text:
            title_found = True
            break
    
    if not title_found:
        raise ValueError("Slide 46 does not match expected title.")
    
    print(f"Verified slide 46 title. Total slides before: {len(prs.slides)}")
    
    # Insert new slide after index 45
    new_slide = insert_slide_after(prs, after_index=45)
    
    # Add title
    title_elem = make_title_shape(
        "LLM-Judged NL Score vs Probability: Subcategory Dissociation"
    )
    new_slide.shapes._spTree.append(title_elem)
    
    # Add image
    new_slide.shapes.add_picture(
        str(IMAGE_PATH),
        left=Emu(501652),
        top=Emu(685800),
        width=Emu(7935422),
        height=Emu(3488161),
    )
    
    # Add annotation bullets
    bullets = [
        "r = \u22120.002 overall: LLM-judged NL consciousness score is completely uncorrelated with baseline probability self-reports.",
        "r = +0.47 for experiential subcategory (p = 0.11): Strongest trend \u2014 models expressing NL uncertainty give higher experiential ratings.",
        "Claude cluster (NL \u2248 47\u201351): moderate baseline probabilities despite unique willingness to express consciousness uncertainty.",
        "Confirms keyword-based finding (r = \u22120.03 overall, r = +0.41 experiential) \u2014 probability paradigm bypasses NL refusal heuristic.",
    ]
    annotation_elem = make_annotation_shape(bullets)
    new_slide.shapes._spTree.append(annotation_elem)
    
    # Save
    prs.save(str(PPTX_PATH))
    
    # Verify
    prs2 = Presentation(str(PPTX_PATH))
    print(f"Total slides after: {len(prs2.slides)}")
    
    for idx in [45, 46, 47]:
        slide = prs2.slides[idx]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.top < 300000:
                print(f"  Slide {idx+1}: \"{shape.text_frame.text[:90]}\"")
                break


if __name__ == "__main__":
    main()
