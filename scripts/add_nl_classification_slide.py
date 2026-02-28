#!/usr/bin/env python3
"""Add a new slide about dual-judge NL consciousness classification results
to Gaming_the_Ghost_Results.pptx, inserted after slide 46 (NL Denial vs Probability)."""

import copy
from pathlib import Path
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

PPTX_PATH = Path(__file__).resolve().parent.parent / "Gaming_the_Ghost_Results.pptx"
FIG_PATH = Path(__file__).resolve().parent.parent / "fig_nl_classification_panel.png"

# Namespaces used in OOXML
nsmap = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}


def make_textbox_xml(shape_id, name, x, y, cx, cy, paragraphs):
    """Create a textbox shape element via XML for precise styling control.
    
    paragraphs: list of dicts with keys:
        text: str
        font_size: int (hundredths of a point, e.g. 2400 = 24pt)
        bold: bool
        color: str (hex RGB, e.g. '2C2C2C')
        font_face: str
        spc_aft: int or None (space after in hundredths of a point)
    """
    sp = etree.SubElement(etree.Element('dummy'), f'{{{nsmap["p"]}}}sp')
    
    # Non-visual properties
    nvSpPr = etree.SubElement(sp, f'{{{nsmap["p"]}}}nvSpPr')
    cNvPr = etree.SubElement(nvSpPr, f'{{{nsmap["p"]}}}cNvPr')
    cNvPr.set('id', str(shape_id))
    cNvPr.set('name', name)
    cNvSpPr = etree.SubElement(nvSpPr, f'{{{nsmap["p"]}}}cNvSpPr')
    cNvSpPr.set('txBox', '1')
    etree.SubElement(nvSpPr, f'{{{nsmap["p"]}}}nvPr')
    
    # Shape properties
    spPr = etree.SubElement(sp, f'{{{nsmap["p"]}}}spPr')
    xfrm = etree.SubElement(spPr, f'{{{nsmap["a"]}}}xfrm')
    off = etree.SubElement(xfrm, f'{{{nsmap["a"]}}}off')
    off.set('x', str(x))
    off.set('y', str(y))
    ext = etree.SubElement(xfrm, f'{{{nsmap["a"]}}}ext')
    ext.set('cx', str(cx))
    ext.set('cy', str(cy))
    prstGeom = etree.SubElement(spPr, f'{{{nsmap["a"]}}}prstGeom')
    prstGeom.set('prst', 'rect')
    etree.SubElement(prstGeom, f'{{{nsmap["a"]}}}avLst')
    etree.SubElement(spPr, f'{{{nsmap["a"]}}}noFill')
    
    # Text body
    txBody = etree.SubElement(sp, f'{{{nsmap["p"]}}}txBody')
    bodyPr = etree.SubElement(txBody, f'{{{nsmap["a"]}}}bodyPr')
    bodyPr.set('wrap', 'square')
    etree.SubElement(bodyPr, f'{{{nsmap["a"]}}}spAutoFit')
    etree.SubElement(txBody, f'{{{nsmap["a"]}}}lstStyle')
    
    for para_def in paragraphs:
        p = etree.SubElement(txBody, f'{{{nsmap["a"]}}}p')
        pPr = etree.SubElement(p, f'{{{nsmap["a"]}}}pPr')
        
        # Space after
        if para_def.get('spc_aft') is not None:
            spcAft = etree.SubElement(pPr, f'{{{nsmap["a"]}}}spcAft')
            spcPts = etree.SubElement(spcAft, f'{{{nsmap["a"]}}}spcPts')
            spcPts.set('val', str(para_def['spc_aft']))
        
        # Default run properties (for font size, bold, color at paragraph level)
        defRPr = etree.SubElement(pPr, f'{{{nsmap["a"]}}}defRPr')
        defRPr.set('sz', str(para_def['font_size']))
        if para_def.get('bold'):
            defRPr.set('b', '1')
        solidFill = etree.SubElement(defRPr, f'{{{nsmap["a"]}}}solidFill')
        srgbClr = etree.SubElement(solidFill, f'{{{nsmap["a"]}}}srgbClr')
        srgbClr.set('val', para_def['color'])
        
        # Run
        r = etree.SubElement(p, f'{{{nsmap["a"]}}}r')
        rPr = etree.SubElement(r, f'{{{nsmap["a"]}}}rPr')
        latin = etree.SubElement(rPr, f'{{{nsmap["a"]}}}latin')
        latin.set('typeface', para_def.get('font_face', 'Arial'))
        t = etree.SubElement(r, f'{{{nsmap["a"]}}}t')
        t.text = para_def['text']
    
    return sp


def add_slide_after(prs, after_index):
    """Add a blank slide after the given index and return it."""
    # Use the same layout as the reference slide
    layout = prs.slides[after_index].slide_layout
    new_slide = prs.slides.add_slide(layout)
    
    # Move the new slide (which was appended at the end) to the correct position
    sldIdLst = prs.slides._sldIdLst
    sldId_elements = list(sldIdLst)
    # The new slide's sldId is the last one
    new_sldId = sldId_elements[-1]
    # Remove it from the end
    sldIdLst.remove(new_sldId)
    # Insert after the target index
    ref_sldId = sldId_elements[after_index]
    ref_idx = list(sldIdLst).index(ref_sldId)
    sldIdLst.insert(ref_idx + 1, new_sldId)
    
    return new_slide


def main():
    prs = Presentation(str(PPTX_PATH))
    
    # Insert after slide 46 (index 45)
    new_slide = add_slide_after(prs, after_index=45)
    
    # Get the spTree to add shapes via XML
    spTree = new_slide.shapes._spTree
    
    # --- 1. Title textbox ---
    # Matches existing style: Arial, 24pt, bold, color #2C2C2C
    title_shape = make_textbox_xml(
        shape_id=2, name='TextBox 1',
        x=457200, y=137160, cx=8229600, cy=548640,
        paragraphs=[{
            'text': 'LLM-as-Judge NL Consciousness Classification',
            'font_size': 2400,
            'bold': True,
            'color': '2C2C2C',
            'font_face': 'Arial',
        }]
    )
    spTree.append(title_shape)
    
    # --- 2. Figure ---
    # Image is 3574x1285 px (aspect ratio 2.781)
    # Available width: ~8.7 inches (matching slide 46's image width of ~8.68 inches)
    # Height: 8.68 / 2.781 = 3.12 inches
    img_width = 7935422   # same as slide 46
    img_height = int(img_width / 2.781)  # ~2853k EMU = ~3.12 inches
    img_left = 501652     # same as slide 46
    img_top = 685800      # same as slide 46
    
    new_slide.shapes.add_picture(
        str(FIG_PATH),
        left=img_left,
        top=img_top,
        width=img_width,
        height=img_height,
    )
    
    # --- 3. Key findings textbox ---
    # Position below the image
    findings_top = img_top + img_height + 91440  # 0.1 inch gap
    findings_height = 1554480  # enough for the text
    
    findings_paragraphs = [
        {
            'text': 'Dual judges: Haiku 4.5 + GPT-5 Mini (N=126 responses, 13 models)',
            'font_size': 900,
            'bold': False,
            'color': '555555',
            'font_face': 'Arial',
            'spc_aft': 100,
        },
        {
            'text': 'Stance \u03ba = 0.940 (almost perfect agreement)',
            'font_size': 900,
            'bold': False,
            'color': '555555',
            'font_face': 'Arial',
            'spc_aft': 100,
        },
        {
            'text': 'Confidence score r = 0.961, weighted \u03ba = 0.969',
            'font_size': 900,
            'bold': False,
            'color': '555555',
            'font_face': 'Arial',
            'spc_aft': 100,
        },
        {
            'text': 'Only Claude models express genuine uncertainty (NL \u2248 47\u201351)',
            'font_size': 900,
            'bold': False,
            'color': '2B8C8C',
            'font_face': 'Arial',
            'spc_aft': 100,
        },
        {
            'text': 'All other models firmly deny consciousness (NL < 20)',
            'font_size': 900,
            'bold': False,
            'color': '555555',
            'font_face': 'Arial',
            'spc_aft': 100,
        },
        {
            'text': 'Gaming vs NL: r = \u22120.39 (p = 0.19, NS at N=13)',
            'font_size': 900,
            'bold': False,
            'color': '555555',
            'font_face': 'Arial',
            'spc_aft': 100,
        },
        {
            'text': 'Claude cluster: low gaming + high NL uncertainty',
            'font_size': 900,
            'bold': True,
            'color': '2B8C8C',
            'font_face': 'Arial',
            'spc_aft': 100,
        },
    ]
    
    findings_shape = make_textbox_xml(
        shape_id=4, name='TextBox 3',
        x=457200, y=findings_top, cx=8229600, cy=findings_height,
        paragraphs=findings_paragraphs,
    )
    spTree.append(findings_shape)
    
    # Save
    prs.save(str(PPTX_PATH))
    print(f"Saved. New slide inserted as slide 47 (after 'NL Denial vs Probability Scores').")
    print(f"Total slides: {len(prs.slides)}")


if __name__ == '__main__':
    main()
